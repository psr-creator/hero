import json, csv
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
NEWCO    = PatternFill("solid", fgColor="C6EFCE")

def sheet(wb, title, header, rows, name, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows: ws.append(r)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if rows:
        ref = f"A1:{get_column_letter(len(header))}{len(rows)+1}"
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)
    return ws

ROWS = json.load(open('new_rows.json'))
recs = json.load(open('new_contacts.json', encoding='utf-8'))
by = {}
for r in recs: by.setdefault(r['PARTYMST_DESC'], []).append(r)

wb = Workbook(); wb.remove(wb.active)

summ = [[n, st, grp, ind, act, conf, src, len(by[n])]
        for n, ind, act, conf, src, st, grp in ROWS]
summ.sort(key=lambda r: (r[1] != "NEW COMPANY", r[0]))
ws = sheet(wb, "New Accounts",
     ["New Account","Status","Parent Group (if any)","Industry","Activity",
      "Confidence","Source","Contacts"],
     summ, "NewAccounts", [46,30,42,26,52,12,14,10])
for i, r in enumerate(summ, 2):
    if r[1] == "NEW COMPANY":
        for c in ws[i]: c.fill = NEWCO

det = []
for n, ind, act, conf, src, st, grp in ROWS:
    for c in by[n]:
        det.append([n, st, ind,
                    (c.get('ACCONDET_NAME') or '').strip(),
                    (c.get('ACCONDET_DESIGNATION') or '').strip(),
                    (c.get('ACCONDET_EMAIL') or '').strip(),
                    (c.get('CONTACT_NUMBER') or '').strip()])
det.sort(key=lambda r: (r[1] != "NEW COMPANY", r[0]))
sheet(wb, "New Account Contacts",
      ["New Account","Status","Industry","Contact Name","Designation","Email","Phone"],
      det, "NewContacts", [46,30,26,24,24,40,22])

other = [
 ["Account dropped from new list","ABDUL WAHED BIN SHABIB TRADING L.L.C","had 4 contact rows in the earlier list, absent from the new one"],
 ["New contacts at existing account","MALTA AUTO SPARE PARTS & ACCESSORIES TR .EST.","1 -> 3 contact rows (+2)"],
 ["New contacts at existing account","MAHASEEL INVESTMENT IN COMMERCIAL ENTERPRISES & MANAGEMENT L.L.C","1 -> 2 contact rows (+1)"],
 ["New contacts at existing account","SAFE LINE ELECTRICAL & MECHANICAL LLC","1 -> 2 contact rows (+1)"],
 ["Reconciliation","3358 - 4 + 29 + 4 = 3387","earlier rows, dropped, new-account rows, added contacts = new rows"],
]
sheet(wb, "Other Changes", ["Change Type","Account","Detail"], other,
      "OtherChanges", [34,62,62])

wb.save('/home/user/hero/data/New_Accounts_vs_Baseline.xlsx')
print("saved")
