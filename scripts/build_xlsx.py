import csv
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)

def read(p):
    r = csv.reader(open(p, encoding='utf-8-sig'))
    return next(r), list(r)

lu_hdr, lu = read('company_industry_lookup.csv')
fl_hdr, fl = read('company_industry_full_list.csv')

wb = Workbook()

def style(ws, hdr, nrows, widths, name):
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(1, c, h)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 28
    for c, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = wdt
    ws.freeze_panes = "A2"
    ref = f"A1:{get_column_letter(len(hdr))}{nrows+1}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)

# --- Sheet 1: summary ---
ws = wb.active; ws.title = "Industry Summary"
cnt_co = Counter(r[1] for r in lu)
cnt_rows = defaultdict(int)
for r in lu:
    cnt_rows[r[1]] += int(r[5])
data = sorted(cnt_co.items(), key=lambda kv: -kv[1])
for i, (ind, n) in enumerate(data, 2):
    ws.cell(i, 1, ind); ws.cell(i, 2, n); ws.cell(i, 3, cnt_rows[ind])
    ws.cell(i, 4, round(n / len(lu) * 100, 1))
style(ws, ["Industry", "Unique companies", "Rows in source list", "% of customer base"],
      len(data), [42, 18, 20, 20], "Summary")

# --- Sheet 2: lookup ---
ws = wb.create_sheet("Company Lookup")
for i, r in enumerate(lu, 2):
    for c, v in enumerate(r, 1):
        ws.cell(i, c, int(v) if c == 6 else v)
style(ws, lu_hdr, len(lu), [58, 32, 56, 13, 15, 18], "Lookup")

# --- Sheet 3: full list in original order ---
ws = wb.create_sheet("Full List (original order)")
for i, r in enumerate(fl, 2):
    for c, v in enumerate(r, 1):
        ws.cell(i, c, v)
style(ws, fl_hdr, len(fl), [58, 32, 56, 13, 15], "FullList")

# --- Sheet 4: needs review ---
ws = wb.create_sheet("Needs Review")
low = [r for r in lu if r[3] == 'Low' or r[1] == 'Unclassified']
low.sort(key=lambda r: (r[1] != 'Unclassified', r[0]))
for i, r in enumerate(low, 2):
    for c, v in enumerate(r, 1):
        ws.cell(i, c, int(v) if c == 6 else v)
style(ws, lu_hdr, len(low), [58, 32, 56, 13, 15, 18], "NeedsReview")

wb.save("UAE_Customer_Industry_Classification.xlsx")
print("sheets:", wb.sheetnames)
print("summary rows:", len(data), "| lookup:", len(lu), "| full:", len(fl), "| needs review:", len(low))
